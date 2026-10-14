#!/usr/bin/env python3
"""Pembaca file tabel (.xlsx / .xls / .csv) untuk impor massal.

Format dideteksi dari ISI file (magic bytes), bukan dari ekstensi, supaya file
yang ekstensinya salah / dihemat ulang sebagai .csv tidak menyebabkan crash.
Semua kegagalan raise ValueError dengan pesan Bahasa Indonesia yang aman
ditampilkan di UI (bukan traceback).

XLSX  -> openpyxl (wajib ada)
XLS   -> xlrd (opsional; kalau belum terpasang muncul pesan jelas)
CSV   -> stdlib csv + deteksi encoding/delimiter
"""
import csv
import io
import itertools
import os
import re
import zipfile

MAX_BYTES = 40 * 1024 * 1024  # batas aman upload (file xl penuh bisa puluhan MB)

XLSX_MAGIC = b"PK\x03\x04"
OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # .xls biner (CFB/Compound File)

MSG_UNKNOWN = ("Format file tidak dikenali. Upload .xlsx (Excel 2007+) atau .csv. "
               "Kalau file aslinya .xls (Excel 97-2003), buka di Excel/LibreOffice "
               "lalu pilih 'Save As' .xlsx atau .csv.")

MSG_ZIP_BUKAN_XLSX = ("File ini zip tapi bukan Excel .xlsx (kemungkinan .ods "
                      "LibreOffice atau .docx). Buka di LibreOffice/Excel lalu "
                      "'Save As' Excel 2007-365 (.xlsx), atau export ke .csv.")


class Sheet:
    """Satu tab/sheet: nama + iterator baris.

    Baris = tuple nilai sel. Pada mode hyperlink, sel adalah objek sel openpyxl.
    """

    def __init__(self, name, rows):
        self.name = name
        self._rows = rows if hasattr(rows, "__next__") else iter(rows)

    @property
    def rows(self):
        return self._rows

    def prepare(self, keywords=("username", "akun"), max_scan=15):
        """Lewati baris judul di awal sheet.

        Return (header, data_iter):
          header  -> list string lowercase sesuai indeks kolom ([] bila tidak ada)
          data_iter -> iterator baris data (sudah melewati baris header)
        """
        buf, header = [], []
        for i, row in enumerate(self._rows):
            if i >= max_scan:
                break
            buf.append(row)
            cells = [str(c).strip().lower() if c is not None else ""
                     for c in row]
            if any(k in c for k in keywords for c in cells):
                header, buf = cells, buf[i + 1:]
                break
        return header, itertools.chain(buf, self._rows)


# ---------- deteksi format ----------

def _looks_text(data):
    return b"\x00" not in data[:8192]


def _zip_kind(data):
    """ZIP apa isi file ini? -> 'xlsx' | 'ods' | 'zip'."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            names = z.namelist()
    except Exception:
        return "zip"
    if any(n.startswith("xl/") for n in names):
        return "xlsx"
    if "content.xml" in names or any("opendocument" in n for n in names):
        return "ods"
    return "zip"


def sniff(data, filename=""):
    """-> 'xlsx' | 'xls' | 'csv' | 'ods' | 'zip' | 'unknown'"""
    if not data:
        return "unknown"
    if data[:4] == XLSX_MAGIC:
        return _zip_kind(data)          # zip: xlsx / ods / lain-lain
    if data[:8] == OLE2_MAGIC:
        return "xls"
    ext = os.path.splitext(filename or "")[1].lower().lstrip(".")
    if ext in ("xlsx", "xlsm"):
        return "unknown"          # Claims xlsx tapi isi bukan zip -> error jelas
    if ext in ("csv", "txt", "tsv"):
        return "csv"
    return "csv" if _looks_text(data) else "unknown"


def decode_text(data):
    """Decode teks dengan deteksi encoding (BOM -> utf-8 -> charset_normalizer)."""
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        try:
            return data.decode("utf-16")
        except UnicodeDecodeError:
            pass
    for enc in ("utf-8-sig",):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    try:
        import charset_normalizer
        best = charset_normalizer.from_bytes(data[:200_000]).best()
        if best is not None:
            return str(best)
    except Exception:
        pass
    return data.decode("latin-1", "replace")


# ---------- pembaca per format ----------

def _xlsx_sheets(data, hyperlinks=False):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError("Modul 'openpyxl' belum terpasang. "
                         "Jalankan: pip install openpyxl")
    try:
        # hyperlink sel hanya terbaca di mode normal (bukan read_only)
        wb = load_workbook(io.BytesIO(data), read_only=not hyperlinks,
                           data_only=True)
    except Exception as e:
        raise ValueError(f"Gagal membaca file .xlsx: {e}")
    sheets = []
    for ws in wb.worksheets:
        if hyperlinks:
            rows = (tuple(c for c in row) for row in ws.iter_rows())
        else:
            rows = (tuple(c for c in row) for row in ws.iter_rows(values_only=True))
        sheets.append(Sheet(ws.title, rows))
    return sheets


def _xls_sheets(data):
    try:
        import xlrd
    except ImportError:
        raise ValueError(
            "File Excel lama (.xls / Excel 97-2003) belum didukung. "
            "Buka di Excel/LibreOffice lalu 'Save As' .xlsx atau .csv, "
            "atau jalankan: pip install xlrd")
    try:
        wb = xlrd.open_workbook(file_contents=data)
    except Exception as e:
        raise ValueError(f"Gagal membaca file .xls: {e}")
    return [Sheet(sh.name, (tuple(sh.row_values(i)) for i in range(sh.nrows)))
            for sh in wb.sheets()]


def _csv_sheet(data, filename=""):
    name = os.path.splitext(os.path.basename(filename or ""))[0] or "CSV"
    text = decode_text(data)
    delim = ","
    try:
        delim = csv.Sniffer().sniff(text[:8192],
                                    delimiters=",;\t|").delimiter
    except csv.Error:
        pass
    return Sheet(name, csv.reader(io.StringIO(text), delimiter=delim))


def read_sheets(data, filename="", hyperlinks=False, max_bytes=MAX_BYTES):
    """Baca file -> list[Sheet]. ValueError bila format/file bermasalah."""
    if not data:
        raise ValueError("File kosong / tidak terbaca.")
    if len(data) > max_bytes:
        raise ValueError(f"File terlalu besar (maks {max_bytes // 1048576} MB).")
    fmt = sniff(data, filename)
    if fmt == "xlsx":
        return _xlsx_sheets(data, hyperlinks)
    if fmt == "xls":
        return _xls_sheets(data)
    if fmt == "csv":
        return [_csv_sheet(data, filename)]
    if fmt in ("ods", "zip"):
        raise ValueError(MSG_ZIP_BUKAN_XLSX)
    raise ValueError(MSG_UNKNOWN)


# ---------- pemetaan kolom / platform ----------

def platform_from_name(name):
    """'INSTAGRAM PEMANTAUAN' -> 'instagram'. None bila tidak dikenali."""
    n = re.sub(r"[^A-Z]", "", (name or "").upper())
    if "INSTAGRAM" in n or n == "IG":
        return "instagram"
    if "TIKTOK" in n:
        return "tiktok"
    if "THREADS" in n or "TRHEADS" in n:  # 'TRHEADS' = salah ketik umum
        return "threads"
    if "FACEBOOK" in n or n == "FB":
        return "facebook"
    return None


def col_index(header, keys, default=None):
    """Index kolom pertama yang header-nya mengandung salah satu keys."""
    for i, h in enumerate(header):
        if h and any(k in h for k in keys):
            return i
    return default


def cell(row, idx):
    """Nilai sel pada indeks idx (None bila indeks di luar baris)."""
    if idx is None or idx < 0 or idx >= len(row):
        return None
    return row[idx]